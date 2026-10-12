"""Regressions for editable Normal body styles and WPS-renamed style IDs."""
import importlib.util
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "convert.py"
SPEC = importlib.util.spec_from_file_location("markdown_to_word", SCRIPT)
CONVERTER = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(CONVERTER)
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


class BodyStyleTests(unittest.TestCase):
    def make_document(self, path):
        doc = Document()
        body = doc.styles["Body Text"]
        body.paragraph_format.space_before = Pt(9)
        body.paragraph_format.space_after = Pt(9)
        first = doc.styles.add_style("First Paragraph", WD_STYLE_TYPE.PARAGRAPH)
        first.base_style = body
        doc.add_heading("标题", level=1)
        for text, style in (("首段 Chinese 123", first), ("后续正文", body), ("已有正文", "Normal")):
            p = doc.add_paragraph(text, style=style)
            p.paragraph_format.first_line_indent = Pt(21)
        doc.add_paragraph("列表内容", style="List Bullet")
        formula = doc.add_paragraph(style=body)
        formula._p.append(parse_xml(
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
            '<m:oMath><m:r><m:t>x</m:t></m:r></m:oMath></m:oMathPara>'
        ))
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].style = body
        cell.paragraphs[0].add_run("表格正文")
        separator = doc.add_paragraph()
        separator._p.get_or_add_pPr().append(parse_xml(
            '<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:bottom w:val="single" w:sz="4"/></w:pBdr>'
        ))
        doc.save(path)

    def read_parts(self, path):
        with ZipFile(path) as archive:
            return {name: archive.read(name) for name in ("word/styles.xml", "word/document.xml")}

    def check_result(self, path, normal_id):
        parts = self.read_parts(path)
        styles = etree.fromstring(parts["word/styles.xml"])
        document = etree.fromstring(parts["word/document.xml"])
        normal = styles.xpath(f'./w:style[@w:styleId="{normal_id}"]', namespaces=NS)[0]
        indent = normal.find("w:pPr/w:ind", NS)
        self.assertEqual(indent.get(qn("w:firstLineChars")), "200")
        self.assertEqual(indent.get(qn("w:firstLine")), "420")
        paragraphs = document.xpath("./w:body/w:p", namespaces=NS)
        for p in paragraphs[1:4]:
            self.assertEqual(p.find("w:pPr/w:pStyle", NS).get(qn("w:val")), normal_id)
            direct = p.find("w:pPr/w:ind", NS)
            self.assertTrue(direct is None or qn("w:firstLineChars") not in direct.attrib)
        for p in paragraphs[1:3]:
            spacing = p.find("w:pPr/w:spacing", NS)
            self.assertEqual(spacing.get(qn("w:before")), "180")
            self.assertEqual(spacing.get(qn("w:after")), "180")
        # Heading and list styles must not acquire the new body indentation.
        for p in (paragraphs[0], paragraphs[4]):
            style_id = p.find("w:pPr/w:pStyle", NS).get(qn("w:val"))
            style = styles.xpath(f'./w:style[@w:styleId="{style_id}"]', namespaces=NS)[0]
            self.assertEqual(style.find("w:pPr/w:ind", NS).get(qn("w:firstLineChars")), "0")
        formula = paragraphs[5]
        self.assertEqual(formula.find("w:pPr/w:pStyle", NS).get(qn("w:val")), normal_id)
        self.assertEqual(formula.find("w:pPr/w:ind", NS).get(qn("w:firstLineChars")), "0")
        cell = document.find("w:body/w:tbl/w:tr/w:tc", NS)
        self.assertEqual(cell.find("w:tcPr/w:vAlign", NS).get(qn("w:val")), "center")
        self.assertEqual(cell.find("w:p/w:pPr/w:jc", NS).get(qn("w:val")), "center")
        self.assertEqual(cell.find("w:p/w:pPr/w:ind", NS).get(qn("w:firstLineChars")), "0")
        self.assertEqual(paragraphs[-1].find("w:pPr/w:ind", NS).get(qn("w:firstLineChars")), "0")
        # The editable body paragraphs have no direct first-line override. Thus
        # reapplying Normal, or inserting a paragraph with that style, yields 2 chars.
        CONVERTER.apply_default_formatting(path)
        self.assertEqual(parts, self.read_parts(path), "Formatting must be idempotent")

    def test_pandoc_body_styles_inherit_normal_indent(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "document.docx"
            self.make_document(path)
            CONVERTER.apply_default_formatting(path)
            self.check_result(path, "Normal")

    def test_wps_numeric_style_ids_are_resolved_by_name(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "document.docx"
            self.make_document(path)
            with ZipFile(path) as archive:
                entries = [(info, archive.read(info)) for info in archive.infolist()]
            ids = {"Normal": "101", "BodyText": "103", "FirstParagraph": "123"}
            rewritten = []
            for info, data in entries:
                if info.filename in ("word/styles.xml", "word/document.xml"):
                    root = etree.fromstring(data)
                    for node in root.iter():
                        for key in (qn("w:styleId"), qn("w:val")):
                            if node.get(key) in ids and etree.QName(node).localname in (
                                "style", "pStyle", "basedOn", "next", "link"
                            ):
                                node.set(key, ids[node.get(key)])
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                rewritten.append((info, data))
            with ZipFile(path, "w") as archive:
                for info, data in rewritten:
                    archive.writestr(info, data)
            CONVERTER.apply_default_formatting(path)
            self.check_result(path, "101")


if __name__ == "__main__":
    unittest.main()
